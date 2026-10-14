"""
多格式作业与量规文件解析服务：
支持：
1. 纯文本 (.txt / .md) -> 本地直接 UTF-8 解码
2. Word 文档 (.docx) -> 使用 python-docx 提取全部段落与表格文字
3. PDF 文档 (.pdf) -> 先用 pypdf 提取文字；若提取文本为空或极少（扫描版 PDF），自动切换为 Gemini 多模态 (Vision) 原生读取
4. 图片文件 (.png / .jpg / .jpeg / .webp) -> 直接交由 Gemini 多模态 (Vision) 原生读取
"""
import io
from pathlib import Path
from typing import Any, Dict, Tuple

from app.services.llm_provider import get_llm_provider


def detect_submission_type_and_mime(filename: str) -> Tuple[str, str]:
    """根据文件名后缀识别作业类型与标准 MIME 类型。"""
    ext = Path(filename).suffix.lower()
    if ext in [".png"]:
        return "image", "image/png"
    if ext in [".jpg", ".jpeg"]:
        return "image", "image/jpeg"
    if ext in [".webp"]:
        return "image", "image/webp"
    if ext == ".pdf":
        return "pdf", "application/pdf"
    if ext == ".docx":
        return (
            "docx",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
    return "text", "text/plain"


def parse_uploaded_file(
    file_bytes: bytes, filename: str, prompt_hint: str = ""
) -> Dict[str, Any]:
    """
    统一解析上传的作业或量规文件：
    返回字典：
    - submission_type: 'text' | 'image' | 'pdf' | 'docx'
    - mime_type: MIME 字符串
    - extracted_text: 解析出的文本内容
    - used_multimodal: 是否使用了 Gemini Vision 多模态能力
    - is_clear: 内容是否清晰可读
    - ambiguity_reason: 若存在扫描模糊或识别受限，记录原因供模糊澄清机制使用
    """
    sub_type, mime_type = detect_submission_type_and_mime(filename)
    llm = get_llm_provider()

    # 1. 纯文本文件 (.txt / .md)
    if sub_type == "text":
        try:
            text = file_bytes.decode("utf-8")
        except UnicodeDecodeError:
            text = file_bytes.decode("gbk", errors="replace")
        return {
            "submission_type": "text",
            "mime_type": mime_type,
            "extracted_text": text.strip(),
            "used_multimodal": False,
            "is_clear": len(text.strip()) >= 20,
            "ambiguity_reason": "" if len(text.strip()) >= 20 else "文本文件内容过短",
        }

    # 2. Word 文档 (.docx)
    if sub_type == "docx":
        try:
            import docx  # type: ignore

            doc = docx.Document(io.BytesIO(file_bytes))
            paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
            # 同时提取 Word 表格中的文字（对量规表格和实验报告尤为关键）
            for table in doc.tables:
                for row in table.rows:
                    cells = [c.text.strip().replace("\n", " ") for c in row.cells if c.text.strip()]
                    if cells:
                        paragraphs.append(" | ".join(cells))
            full_text = "\n".join(paragraphs).strip()
            return {
                "submission_type": "docx",
                "mime_type": mime_type,
                "extracted_text": full_text,
                "used_multimodal": False,
                "is_clear": len(full_text) >= 20,
                "ambiguity_reason": "" if len(full_text) >= 20 else "Word 文档中未提取到有效文字内容",
            }
        except Exception as exc:
            return {
                "submission_type": "docx",
                "mime_type": mime_type,
                "extracted_text": f"[Word 文档解析异常: {exc}]",
                "used_multimodal": False,
                "is_clear": False,
                "ambiguity_reason": f"Word 文档无法正常解析：{exc}",
            }

    # 3. PDF 文档 (.pdf)：优先尝试本地文本层提取，若为扫描版 PDF 则调用 Gemini 多模态原生读取
    if sub_type == "pdf":
        extracted_pages = []
        try:
            from pypdf import PdfReader  # type: ignore

            reader = PdfReader(io.BytesIO(file_bytes))
            for page in reader.pages:
                page_txt = (page.extract_text() or "").strip()
                if page_txt:
                    extracted_pages.append(page_txt)
        except Exception:
            pass

        local_pdf_text = "\n\n".join(extracted_pages).strip()
        if len(local_pdf_text) >= 35:
            return {
                "submission_type": "pdf",
                "mime_type": mime_type,
                "extracted_text": local_pdf_text,
                "used_multimodal": False,
                "is_clear": True,
                "ambiguity_reason": "",
            }

        # 本地未能提取出足够文本 -> 判定为扫描版 PDF，直接调用 Gemini 多模态原生读取
        vision_res = llm.extract_multimodal_content(
            file_bytes=file_bytes, mime_type="application/pdf", prompt_hint=prompt_hint
        )
        return {
            "submission_type": "pdf",
            "mime_type": mime_type,
            "extracted_text": vision_res.get("extracted_text", ""),
            "used_multimodal": True,
            "is_clear": bool(vision_res.get("is_clear", True)),
            "ambiguity_reason": vision_res.get("ambiguity_reason", ""),
        }

    # 4. 图片作业 (.png / .jpg / .webp)：直接利用 Gemini 多模态 (Vision) 原生读取
    vision_res = llm.extract_multimodal_content(
        file_bytes=file_bytes, mime_type=mime_type, prompt_hint=prompt_hint
    )
    return {
        "submission_type": "image",
        "mime_type": mime_type,
        "extracted_text": vision_res.get("extracted_text", ""),
        "used_multimodal": True,
        "is_clear": bool(vision_res.get("is_clear", True)),
        "ambiguity_reason": vision_res.get("ambiguity_reason", ""),
    }
